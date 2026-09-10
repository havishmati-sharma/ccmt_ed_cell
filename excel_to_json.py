"""Convert an Excel question bank (.xlsx) to questions.json.

Usage:
    python excel_to_json.py
    python excel_to_json.py --excel questions.xlsx --json questions.json
    python excel_to_json.py --excel questions.xlsx --json questions.json --sheet Sheet1
    python excel_to_json.py --excel questions.xlsx --json questions.json --title "My Quiz"

Install the one-time dependency:
    python -m pip install openpyxl
"""

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

DEFAULT_INPUT_FILE = "questions.xlsx"
DEFAULT_OUTPUT_FILE = "questions.json"

DEFAULT_TITLE = (
    "ICSE Class 10 Computer Applications - "
    "Strings, Arrays and User-defined Methods"
)

REQUIRED_COLUMNS = [
    "id",
    "subject",
    "question",
    "option_a",
    "option_b",
    "option_c",
    "option_d",
    "correct_option",
]


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def convert_excel_to_json(input_file, output_file, sheet_name=None, title=None):
    workbook = load_workbook(
        filename=input_file,
        read_only=True,
        data_only=True
    )

    if sheet_name is None:
        worksheet = workbook.active
    elif sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
    else:
        raise ValueError(
            "Sheet not found: " + sheet_name
            + ". Available sheets: "
            + ", ".join(workbook.sheetnames)
        )

    rows = worksheet.iter_rows(values_only=True)

    try:
        first_row = next(rows)
    except StopIteration:
        raise ValueError("The Excel sheet is empty.")

    headers = [clean(value).lower() for value in first_row]
    missing = [column for column in REQUIRED_COLUMNS if column not in headers]

    if missing:
        raise ValueError("Missing columns: " + ", ".join(missing))

    questions = []
    used_ids = set()

    for excel_row_number, values in enumerate(rows, start=2):
        if all(value is None or clean(value) == "" for value in values):
            continue

        row = dict(zip(headers, values))

        try:
            question_id = int(row["id"])
        except (TypeError, ValueError):
            raise ValueError(
                "Invalid id at Excel row " + str(excel_row_number)
            )

        if question_id in used_ids:
            raise ValueError(
                "Duplicate id " + str(question_id)
                + " at Excel row " + str(excel_row_number)
            )
        used_ids.add(question_id)

        correct_option = clean(row["correct_option"]).lower()
        if correct_option not in {"a", "b", "c", "d"}:
            raise ValueError(
                "correct_option at Excel row "
                + str(excel_row_number)
                + " must be a, b, c, or d."
            )

        # QuestionBankQuiz.java expects "options" as a flat array of answer
        # text, and "correctAnswer" holding the actual text of the correct
        # option (not its letter) - the letter/id is only used in Excel to
        # make the spreadsheet easier to fill in.
        options = []
        correct_answer_text = None

        for option_id in ["a", "b", "c", "d"]:
            option_text = clean(row["option_" + option_id])
            if not option_text:
                raise ValueError(
                    "option_" + option_id
                    + " is empty at Excel row "
                    + str(excel_row_number)
                )
            options.append(option_text)
            if option_id == correct_option:
                correct_answer_text = option_text

        questions.append({
            "id": question_id,
            "subject": clean(row["subject"]),
            "question": clean(row["question"]),
            "options": options,
            "correctAnswer": correct_answer_text,
            "explanation": clean(row.get("explanation", "")),
        })

    result = {
        "version": 1,
        "title": title if title else DEFAULT_TITLE,
        "questions": questions,
    }

    Path(output_file).write_text(
        json.dumps(result, indent=2, ensure_ascii=False),
        encoding="utf-8"
    )

    print("Converted questions:", len(questions))
    print("Created:", output_file)


def main():
    parser = argparse.ArgumentParser(
        description="Convert an Excel question bank (.xlsx) to questions.json."
    )
    parser.add_argument(
        "--excel", "-e",
        dest="input_file",
        default=DEFAULT_INPUT_FILE,
        help="Path to the source .xlsx file (default: " + DEFAULT_INPUT_FILE + ")",
    )
    parser.add_argument(
        "--json", "-j",
        dest="output_file",
        default=DEFAULT_OUTPUT_FILE,
        help="Path to write the generated JSON to (default: " + DEFAULT_OUTPUT_FILE + ")",
    )
    parser.add_argument(
        "--sheet",
        dest="sheet_name",
        default=None,
        help="Worksheet name to read (default: the active sheet)",
    )
    parser.add_argument(
        "--title",
        dest="title",
        default=None,
        help="Quiz title stored in the JSON file (default: a generic ICSE title)",
    )

    args = parser.parse_args()

    try:
        convert_excel_to_json(
            args.input_file, args.output_file, args.sheet_name, args.title
        )
    except Exception as error:
        print("Conversion failed:", error)
        raise SystemExit(1)


if __name__ == "__main__":
    main()
